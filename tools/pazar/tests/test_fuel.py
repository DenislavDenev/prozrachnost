"""The weekly Oil Bulletin of the European Commission. The fixture is the real history file cut to its 6 newest weeks."""
import datetime as dt
import io
from pathlib import Path

import openpyxl
import pytest

from ingest import checks, config, fuel
from ingest.parse import ShapeError

RAW = (Path(__file__).parent / "fixtures" / "oil-bulletin-history-cut.xlsx").read_bytes()
NOW = dt.datetime(2026, 10, 2, 12, 0, tzinfo=dt.timezone.utc)


def test_the_parser_reads_bulgaria_the_union_and_the_euro_area_by_header_name():
    rows = fuel.parse(RAW)
    assert {r[1] for r in rows} == {"BG", "EU", "EUR"} and {r[2] for r in rows} == {"euro95", "diesel", "LPG"}
    by = {(r[0], r[1], r[2]): (r[3], r[4]) for r in rows}
    # the file: Bulgaria, week of 28.09.2026, with taxes: Euro-super 95 1674.8, diesel 1927.7, LPG 682.8 (EUR per 1000 l)
    assert by[(dt.date(2026, 9, 28), "BG", "euro95")][0] == pytest.approx(1674.8, abs=0.05)
    assert by[(dt.date(2026, 9, 28), "BG", "diesel")][0] == pytest.approx(1927.7, abs=0.05)
    assert by[(dt.date(2026, 9, 28), "BG", "LPG")][0] == pytest.approx(682.8, abs=0.05)
    assert by[(dt.date(2026, 9, 28), "BG", "euro95")][1] == pytest.approx(1032.65, abs=0.05)      # without taxes
    assert len({r[0] for r in rows}) == 6


def altered(fn):
    wb = openpyxl.load_workbook(io.BytesIO(RAW))
    fn(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_a_missing_column_or_sheet_or_a_bad_price_stops_the_import():
    def rename(wb):
        ws = wb["Prices with taxes"]
        for c in ws[1]:
            if c.value == "BG_price_with_tax_diesel":
                c.value = "BG_price_with_tax_gasoil"
    with pytest.raises(ShapeError, match="Липсва колона"):
        fuel.parse(altered(rename))
    with pytest.raises(ShapeError, match="Липсва лист"):
        fuel.parse(altered(lambda wb: wb.remove(wb["Prices wo taxes"])))

    def negative(wb):
        ws = wb["Prices with taxes"]
        col = [c.value for c in ws[1]].index("BG_price_with_tax_euro95") + 1
        ws.cell(row=4, column=col).value = -5
    with pytest.raises(ShapeError, match="Невалидна цена"):
        fuel.parse(altered(negative))
    with pytest.raises(ShapeError):
        fuel.parse(b"<html>404</html>")


def test_store_is_idempotent_and_keeps_the_raw_file(c, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA", tmp_path)
    r = fuel.store(c, RAW, now=NOW)
    assert r["state"] == "stored" and r["new"] == 6 * 3 * 3 and r["newest"] == "2026-09-28"
    again = fuel.store(c, RAW, now=NOW)
    assert again["new"] == 0 and again["changed"] == 0
    assert c.execute("SELECT count(*) FROM ops.change_log").fetchone()[0] == 0
    assert len(list((tmp_path / "raw" / "oil-bulletin").glob("*.xlsx"))) == 1
    assert c.execute("SELECT count(*) FROM gold.fuel_week WHERE geo = 'BG' AND fuel = 'euro95' AND week = '2026-09-28' AND with_tax_eur_l = 1.6748").fetchone()[0] == 1


def test_a_changed_value_is_logged_and_a_smaller_answer_is_held(c, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA", tmp_path)
    fuel.store(c, RAW, now=NOW)

    def bump(wb):
        ws = wb["Prices with taxes"]
        col = [x.value for x in ws[1]].index("BG_price_with_tax_diesel") + 1
        ws.cell(row=5, column=col).value = 2000.0
    r = fuel.store(c, altered(bump), now=NOW + dt.timedelta(days=7))
    assert r["changed"] == 1
    assert c.execute("SELECT cause, ref FROM ops.change_log").fetchone() == ("rewritten", "2026-09-21/BG/diesel")

    def cut(wb):
        for ws in wb.worksheets:
            ws.delete_rows(5, 3)          # three weeks fewer
    small = altered(cut)
    assert fuel.store(c, small, now=NOW + dt.timedelta(days=8))["state"] == "held"
    assert fuel.store(c, small, now=NOW + dt.timedelta(days=8, hours=2))["state"] == "held"
    assert c.execute("SELECT count(DISTINCT week) FROM silver.fuel_week").fetchone()[0] == 6        # nothing was replaced
    assert fuel.store(c, small, now=NOW + dt.timedelta(days=9, hours=1))["state"] == "stored"      # the same answer a day later
    assert c.execute("SELECT count(*) FROM ops.held").fetchone()[0] == 0


def test_freshness_wants_a_recent_read_and_a_recent_week(c, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA", tmp_path)
    assert any("горивата" in p for p in fuel_problems(c))
    fuel.store(c, RAW, now=NOW)
    assert not any("горивата" in p for p in fuel_problems(c, NOW))
    assert any("горивата" in p for p in fuel_problems(c, NOW + dt.timedelta(days=12)))


def fuel_problems(c, now=NOW):
    # freshness needs a built day to talk about; the fuel part is checked on its own
    c.execute("INSERT INTO silver.day (day, status, zip_sha256, zip_path, zip_bytes) VALUES ('2026-10-01', 'built', 'x', 'x', 1) ON CONFLICT DO NOTHING")
    return checks.freshness(c, today=now.date(), now=now, st={"seen": {"2026-10-01": {}}, "last_ok": "2026-10-02T05:40:00Z"})
