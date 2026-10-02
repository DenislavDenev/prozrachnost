"""Fuel prices: the weekly Oil Bulletin of the European Commission (a published XLSX, history from 2005).

One request per run. The raw file is kept by sha256; a smaller file (fewer weeks, a value that vanished) is held until
a second read at least a day later gives the same file. Prices are in EUR per 1000 litres, with and without taxes.
The NRA files of average prices from fiscal receipts are not read here: see docs/sources.md.
"""
import datetime as dt
import hashlib
import io
import urllib.request

from . import config
from .parse import ShapeError

URL = "https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
SHEETS = {"Prices with taxes": "with_tax", "Prices wo taxes": "wo_tax"}
GEOS = ("BG", "EU", "EUR")          # Bulgaria, the EU average, the euro area average
FUELS = ("euro95", "diesel", "LPG")


def fetch(url=URL):
    req = urllib.request.Request(url, headers={"User-Agent": config.USER_AGENT})
    with urllib.request.urlopen(req, timeout=120) as r:
        ctype = r.headers.get("Content-Type", "")
        raw = r.read()
    if "spreadsheetml" not in ctype or raw[:2] != b"PK":
        raise ShapeError(f"Очакван XLSX, получено {ctype or 'нищо'} ({len(raw)} байта)")
    return raw


def parse(raw):
    """[(week, geo, fuel, with_tax, wo_tax)] for the geos and fuels above. Columns are found by header name."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise ShapeError("XLSX не се чете: " + str(e)) from e
    got = {}
    for sheet, kind in SHEETS.items():
        if sheet not in wb.sheetnames:
            raise ShapeError("Липсва лист " + sheet)
        rows = ws_rows(wb[sheet])
        header = rows[0]
        cols = {}
        for geo in GEOS:
            for fuel in FUELS:
                name = f"{geo}_price_with_tax_{fuel}" if kind == "with_tax" else f"{geo}_price_wo_tax_{fuel}"
                if name not in header:
                    raise ShapeError("Липсва колона " + name)
                cols[(geo, fuel)] = header.index(name)
        if str(rows[2][0]).strip() != "Date":
            raise ShapeError("Променена структура: очаквано `Date` в третия ред")
        for r in rows[3:]:
            week = r[0]
            if not isinstance(week, dt.datetime):
                continue
            for (geo, fuel), i in cols.items():
                v = r[i] if i < len(r) else None
                if v is None or v == "":
                    continue
                if not isinstance(v, (int, float)) or v <= 0:
                    raise ShapeError(f"Невалидна цена {v!r} за {week:%Y-%m-%d} {geo} {fuel}")
                got.setdefault((week.date(), geo, fuel), {})[kind] = round(float(v), 4)
    if not got:
        raise ShapeError("Празен бюлетин")
    return sorted((w, g, f, v.get("with_tax"), v.get("wo_tax")) for (w, g, f), v in got.items())


def ws_rows(ws):
    return [list(r) for r in ws.iter_rows(min_row=1, max_row=None, values_only=True)]


def store(c, raw, now=None):
    """Raw file by sha256, silver.fuel_week upserted, changes in ops.change_log. Returns the report."""
    now = now or dt.datetime.now(dt.timezone.utc)
    sha = hashlib.sha256(raw).hexdigest()
    rows = parse(raw)
    path = config.DATA / "raw" / "oil-bulletin" / f"{now:%Y-%m-%d}.{sha[:12]}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        path.write_bytes(raw)
    c.execute("INSERT INTO ops.raw_file (source, ref, path, sha256, bytes) VALUES ('oil-bulletin', 'history', %s, %s, %s) ON CONFLICT DO NOTHING",
              (str(path), sha, len(raw)))
    have = {(w, g, f): (a, b) for w, g, f, a, b in c.execute("SELECT week, geo, fuel, with_tax, wo_tax FROM silver.fuel_week")}
    new = {(w, g, f): (a, b) for w, g, f, a, b in rows}
    gone = [k for k in have if k not in new]
    changed = [k for k in new if k in have and tuple(map(_f, have[k])) != tuple(map(_f, new[k]))]
    if (gone or len(new) < len(have)) and have:
        held = c.execute("SELECT sha256, first_seen FROM ops.held WHERE ref = 'oil-bulletin'").fetchone()
        if not (held and held[0] == sha and now - held[1] >= dt.timedelta(days=1)):
            if not held or held[0] != sha:
                c.execute("""INSERT INTO ops.held (ref, sha256, first_seen, reason) VALUES ('oil-bulletin', %s, %s, %s)
                             ON CONFLICT (ref) DO UPDATE SET sha256 = EXCLUDED.sha256, first_seen = EXCLUDED.first_seen, reason = EXCLUDED.reason""",
                          (sha, now, f"по-малко стойности: {len(new)} срещу {len(have)}, изчезнали {len(gone)}"))
                c.execute("INSERT INTO ops.change_log (source, ref, field, old, new, cause) VALUES ('oil-bulletin', 'history', NULL, NULL, %s, 'held')", (sha,))
            return {"state": "held", "rows": len(have), "problems": []}
    with c.transaction():
        for k in changed:
            c.execute("INSERT INTO ops.change_log (source, ref, field, old, new, cause) VALUES ('oil-bulletin', %s, 'price', %s, %s, 'rewritten')",
                      (f"{k[0]}/{k[1]}/{k[2]}", str(have[k]), str(new[k])))
        added = [k for k in new if k not in have]
        with c.cursor() as cur:
            cur.executemany("""INSERT INTO silver.fuel_week (week, geo, fuel, with_tax, wo_tax, source_sha256) VALUES (%s, %s, %s, %s, %s, %s)
                               ON CONFLICT (week, geo, fuel) DO UPDATE SET with_tax = EXCLUDED.with_tax, wo_tax = EXCLUDED.wo_tax, source_sha256 = EXCLUDED.source_sha256""",
                            [(w, g, f, a, b, sha) for (w, g, f), (a, b) in new.items()])
        c.execute("DELETE FROM ops.held WHERE ref = 'oil-bulletin'")
        c.execute("""INSERT INTO silver.fuel_read (read_at, sha256, weeks, newest) VALUES (%s, %s, %s, %s)""",
                  (now, sha, len({k[0] for k in new}), max(k[0] for k in new)))
    return {"state": "stored", "new": len(added), "changed": len(changed), "newest": str(max(k[0] for k in new)), "problems": []}


def _f(v):
    return None if v is None else round(float(v), 4)
